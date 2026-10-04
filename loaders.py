"""Step 1 - Document ingestion.
Every loader returns a list of sections: {"heading": str, "text": str}.
Keeping the heading separate lets the chunker do section-aware chunking later."""
import os
import re

PDF_HEADING = re.compile(r"^\d+\.\s+(\S.{1,60})$")   # "1. Annual Leave"


def _clean(s):
    return re.sub(r"\s+", " ", s).strip()


def load_txt(path):
    """TXT convention: a line starting with '# ' is a heading. Text before the first heading is the title."""
    sections, head, body = [], None, []
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.rstrip("\n")
            if line.startswith("# "):
                if head is not None:
                    sections.append({"heading": head, "text": _clean(" ".join(body))})
                head, body = line[2:].strip(), []
            elif head is not None:
                body.append(line)
    if head is not None:
        sections.append({"heading": head, "text": _clean(" ".join(body))})
    return sections


def load_pdf(path):
    """PDFs carry no style info, so headings are detected by the numbered pattern '1. Title'."""
    from pypdf import PdfReader
    text = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    sections, head, body = [], None, []
    for line in text.splitlines():
        line = line.strip()
        m = PDF_HEADING.match(line)
        if m and not line.endswith("."):
            if head is not None:
                sections.append({"heading": head, "text": _clean(" ".join(body))})
            head, body = m.group(1), []
        elif head is not None and line:
            body.append(line)
    if head is not None:
        sections.append({"heading": head, "text": _clean(" ".join(body))})
    return sections


def load_docx(path):
    """Uses Word paragraph styles: 'Heading N' starts a section, 'Title' is skipped."""
    import docx
    sections, head, body = [], None, []
    for p in docx.Document(path).paragraphs:
        style = p.style.name
        if style.startswith("Heading"):
            if head is not None:
                sections.append({"heading": head, "text": _clean(" ".join(body))})
            head, body = p.text.strip(), []
        elif style != "Title" and head is not None and p.text.strip():
            body.append(p.text)
    if head is not None:
        sections.append({"heading": head, "text": _clean(" ".join(body))})
    return sections


def load_pptx(path):
    """One slide = one section: slide title is the heading, bullets become sentences."""
    from pptx import Presentation
    sections = []
    for slide in Presentation(path).slides:
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        bullets = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        bullets.append(t if t[-1] in ".!?" else t + ".")
        sections.append({"heading": title, "text": _clean(" ".join(bullets))})
    return sections


LOADERS = {".txt": load_txt, ".pdf": load_pdf, ".docx": load_docx, ".pptx": load_pptx}


def load_kb(kb_dir, verbose=True):
    """Load every supported file in kb_dir -> flat list of {source, heading, text}."""
    docs = []
    for name in sorted(os.listdir(kb_dir)):
        ext = os.path.splitext(name)[1].lower()
        if ext not in LOADERS:
            continue
        secs = LOADERS[ext](os.path.join(kb_dir, name))
        chars = sum(len(s["text"]) for s in secs)
        if verbose:
            print(f"  loaded {name:<26} {len(secs):>2} sections {chars:>6} chars")
        if chars == 0:
            print(f"  WARNING: no text extracted from {name} (scanned PDF or empty file?)")
        docs += [{"source": name, **s} for s in secs]
    return docs


if __name__ == "__main__":
    load_kb(os.path.join(os.path.dirname(os.path.abspath(__file__)), "kb"))
