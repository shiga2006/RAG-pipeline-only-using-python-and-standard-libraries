"""Generate the sample knowledge base (PDF, DOCX, TXT, PPTX) into ./kb from kb_content.py."""
import os
from kb_content import DOCS

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "kb")


def write_txt(path, title, sections):
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"{title}\n\n")
        for h, t in sections:
            f.write(f"# {h}\n{t}\n\n")


def write_docx(path, title, sections):
    import docx
    d = docx.Document()
    d.add_heading(title, level=0)
    for h, t in sections:
        d.add_heading(h, level=1)
        d.add_paragraph(t)
    d.save(path)


def write_pdf(path, title, sections):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    st = getSampleStyleSheet()
    story = [Paragraph(title, st["Title"]), Spacer(1, 12)]
    for i, (h, t) in enumerate(sections, 1):
        story += [Paragraph(f"{i}. {h}", st["Heading2"]), Paragraph(t, st["Normal"]), Spacer(1, 8)]
    SimpleDocTemplate(path, pagesize=A4).build(story)


def write_pptx(path, title, slides):
    from pptx import Presentation
    prs = Presentation()
    for h, bullets in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1])      # Title + Content
        s.shapes.title.text = h
        tf = s.placeholders[1].text_frame
        tf.text = bullets[0]
        for b in bullets[1:]:
            tf.add_paragraph().text = b
    prs.save(path)


WRITERS = {".txt": write_txt, ".docx": write_docx, ".pdf": write_pdf, ".pptx": write_pptx}

if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    for name, (title, sections) in DOCS.items():
        WRITERS[os.path.splitext(name)[1]](os.path.join(OUT, name), title, sections)
        print("created", name)
